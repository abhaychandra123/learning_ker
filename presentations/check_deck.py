"""Structural/text-fit checks and approximate layout previews (not Office rendering)."""
from pathlib import Path
from io import BytesIO
from pptx import Presentation
from PIL import Image,ImageDraw,ImageFont
R=Presentation(Path(__file__).parent/'ISTA_FISTA_LISTA_Professor.pptx')
OUT=Path(__file__).parent/'previews';OUT.mkdir(exist_ok=True)
S=96/914400
fonts=Path('/System/Library/Fonts/Supplemental')
def col(c,default='18293B'):
 try:return '#'+str(c.rgb)
 except:return '#'+default
def text(draw,tf,rect,label):
 if not tf.text.strip():return
 x,y,w,h=rect;x+=tf.margin_left*S;y+=tf.margin_top*S;w-=(tf.margin_left+tf.margin_right)*S;h-=(tf.margin_top+tf.margin_bottom)*S
 start=y
 for ip,p in enumerate(tf.paragraphs):
  run=p.runs[0] if p.runs else None
  font=p.font
  sz=font.size or (run.font.size if run else None)
  sz=sz.pt if sz else 18
  fn=ImageFont.truetype(str(fonts/('Arial Bold.ttf' if font.bold else 'Arial.ttf')),round(sz*96/72))
  color=col(font.color)
  lines=[]
  for raw in p.text.split('\n'):
   line=''
   for word in raw.split(' '):
    cand=(line+' '+word).strip()
    if draw.textlength(cand,font=fn)>w and line:lines.append(line);line=word
    else:line=cand
   lines.append(line)
  for line in lines:
   xx=x
   if p.alignment and int(p.alignment)==3:xx=x+w-draw.textlength(line,font=fn)
   elif p.alignment and int(p.alignment)==2:xx=x+(w-draw.textlength(line,font=fn))/2
   draw.text((xx,y),line,fill=color,font=fn);y+=sz*96/72*1.15
  if ip<len(tf.paragraphs)-1:y+=(p.space_after.pt if p.space_after else 0)*96/72
 if y-start>h+3:print('POSSIBLE TEXT OVERFLOW',label,round(y-start),round(h),tf.text[:90])
for i,s in enumerate(R.slides,1):
 im=Image.new('RGB',(1280,720),'white');d=ImageDraw.Draw(im)
 for sh in s.shapes:
  rect=(sh.left*S,sh.top*S,sh.width*S,sh.height*S);x,y,w,h=rect
  if sh.shape_type==13:
   pic=Image.open(BytesIO(sh.image.blob)).convert('RGBA');iw,ih=pic.size
   pic=pic.crop((round(sh.crop_left*iw),round(sh.crop_top*ih),round((1-sh.crop_right)*iw),round((1-sh.crop_bottom)*ih))).resize((round(w),round(h)),Image.Resampling.LANCZOS)
   im.paste(pic,(round(x),round(y)),pic)
  elif sh.has_table:
   yy=y
   for ri,row in enumerate(sh.table.rows):
    xx=x
    for ci,cell in enumerate(row.cells):
     ww=sh.table.columns[ci].width*S;hh=row.height*S
     d.rectangle((xx,yy,xx+ww,yy+hh),fill=col(cell.fill.fore_color,'FFFFFF'))
     text(d,cell.text_frame,(xx,yy,ww,hh),f'{i}:table{ri},{ci}');xx+=ww
    yy+=row.height*S
  else:
   try:
    if sh.fill.type==1:d.rectangle((x,y,x+w,y+h),fill=col(sh.fill.fore_color,'FFFFFF'))
   except:pass
   if sh.has_text_frame:text(d,sh.text_frame,rect,f'{i}:{sh.name}')
 im.save(OUT/f'slide_{i:02d}.png')
for start in range(1,len(R.slides)+1,8):
 canvas=Image.new('RGB',(1280,4*200),(210,214,220));dd=ImageDraw.Draw(canvas)
 for j,i in enumerate(range(start,min(start+8,len(R.slides)+1))):
  im=Image.open(OUT/f'slide_{i:02d}.png');im.thumbnail((620,349))
  # compact 4 columns, 2 rows per montage
  im.thumbnail((310,175));xx=(j%4)*320;yy=(j//4)*200;canvas.paste(im,(xx,yy));dd.text((xx+4,yy+178),str(i),fill='black')
 canvas.crop((0,0,1280,400)).save(OUT/f'montage_{start:02d}.png')
print('Approximate previews:',OUT)
