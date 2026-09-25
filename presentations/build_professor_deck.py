"""Rebuild the professor presentation entirely from the two saved result JSONs."""
import os
os.environ.setdefault('MPLCONFIGDIR','/tmp/learning_ker_mpl')
import json, hashlib
from pathlib import Path
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'presentations'; AS=OUT/'assets'; AS.mkdir(exist_ok=True)
DATA=ROOT/'notebooks/runs_results'
M=[json.loads((DATA/n).read_text()) for n in ['metrics.json','metrics-2.json']]
assert M[0]['source_sha256']==M[1]['source_sha256']
assert M[0]['input_sha256']==M[1]['input_sha256']
NAVY='162C46'; BLUE='3478B6'; ORANGE='D9782F'; GREEN='30866A'; GRAY='526174'; LIGHT='EDF2F6'; WHITE='FFFFFF'; DARK='18293B'
COL={'ista':'#'+BLUE,'lista':'#'+ORANGE,'fista':'#'+GREEN}
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':14,'axes.spines.top':False,'axes.spines.right':False,'pdf.fonttype':42})
def scores(i,key,metric='nrmse_range'):
 return np.array([r[metric] for r in M[i]['comparisons'][key]['metrics']])
def chart_save(fig,name):
 p=AS/(name+'.png');fig.savefig(p,dpi=200,bbox_inches='tight',facecolor='white');plt.close(fig);return p
# Matched depth charts, with exact metrics.
fig,axs=plt.subplots(1,2,figsize=(12.5,4),sharey=True,layout='constrained')
for i,ax in enumerate(axs):
 d=M[i]['config']['layers'];x=np.arange(3)
 for shift,key,lab in [(-.25,f'ista_{d}','ISTA'),(0,f'fista_{d}','FISTA'),(.25,'lista','LISTA')]:
  bars=ax.bar(x+shift,100*scores(i,key),.24,color=COL[lab.lower()],label=lab)
  ax.bar_label(bars,fmt='%.2f',padding=3,fontsize=11)
 ax.set_xticks(x,['Bx','By','Bz']);ax.set_ylim(0,4.8);ax.set_title(f'Run {i+1}: {d} updates / layers');ax.legend(ncol=3,fontsize=10);ax.grid(axis='y',alpha=.15)
axs[0].set_ylabel('Range NRMSE (%)');chart_save(fig,'equal_depth')
fig,axs=plt.subplots(1,3,figsize=(12.5,4),sharey=True,layout='constrained')
for ax,(name,k1,k2) in zip(axs,[('LISTA','lista','lista'),('ISTA','ista_200','ista_500'),('FISTA','fista_200','fista_500')]):
 for shift,i,key,label in [(-.18,0,k1,'8 layers' if name=='LISTA' else '200 iterations'),(.18,1,k2,'16 layers' if name=='LISTA' else '500 iterations')]:
  bars=ax.bar(np.arange(3)+shift,100*scores(i,key),.34,label=label,color=['#3478B6','#D9782F'][i]);ax.bar_label(bars,fmt='%.2f',fontsize=10,padding=3)
 ax.set_xticks(range(3),['Bx','By','Bz']);ax.set_title(name);ax.set_ylim(0,3.6);ax.legend(fontsize=10);ax.grid(axis='y',alpha=.15)
axs[0].set_ylabel('Range NRMSE (%)');chart_save(fig,'between_runs')
fig,axs=plt.subplots(1,3,figsize=(12.5,4),layout='constrained')
for c,ax in enumerate(axs):
 for method in ['ista','fista']:
  h=M[1]['histories'][method];ax.plot([r['iteration'] for r in h],[100*r['components'][c]['nrmse_range'] for r in h],label=method.upper(),color=COL[method],lw=2)
 h=M[1]['histories']['lista'];ax.plot([r['iteration'] for r in h],[100*r['components'][c]['nrmse_range'] for r in h],label='LISTA (16 layers)',color=COL['lista'],lw=2)
 best=min(M[1]['histories']['fista'],key=lambda r:r['components'][c]['nrmse_range'])
 ax.scatter(best['iteration'],100*best['components'][c]['nrmse_range'],color=COL['fista'],s=35,zorder=5)
 ax.annotate(f"Best recorded FISTA\nk={best['iteration']}: {100*best['components'][c]['nrmse_range']:.3f}%",(best['iteration'],100*best['components'][c]['nrmse_range']),xytext=(25,35),textcoords='offset points',fontsize=10,arrowprops={'arrowstyle':'-','color':'gray'})
 ax.set_title(['Bx','By','Bz'][c]);ax.set_xlabel('Iterations / trained layers');ax.grid(alpha=.15)
axs[0].set_ylabel('Range NRMSE (%)');axs[2].legend(fontsize=9);chart_save(fig,'reference_history')
fig,axs=plt.subplots(1,2,figsize=(12,4),layout='constrained')
for method in ['ista','fista','lista']:
 h=M[1]['histories'][method]
 for ax,metric in zip(axs,['objective','pg_rms_uT']):
  ax.semilogy([r['iteration'] for r in h],[r['components'][0][metric] for r in h],color=COL[method],label=method.upper(),lw=2)
for ax,title in zip(axs,['Bx: original L1 objective','Bx: stationarity RMS (µT)']):ax.set_title(title);ax.set_xlabel('Iterations / trained layers');ax.grid(alpha=.15);ax.legend(fontsize=11)
chart_save(fig,'optimization_history')
fig,axs=plt.subplots(1,2,figsize=(12,4),layout='constrained')
for i in range(2):
 h=M[i]['training'];axs[0].plot([r['epoch'] for r in h],[r['training_mse_normalized'] for r in h],label=f"{M[i]['config']['layers']} layers",lw=2)
 axs[1].plot([r['epoch'] for r in h],[r['threshold_normalized']*M[i]['scale_uT'] for r in h],label=f"{M[i]['config']['layers']} layers",lw=2)
for ax,title in zip(axs,['Final-output training MSE (normalized)','Learned threshold in physical units (µT)']):ax.set_title(title);ax.set_xlabel('Training epoch');ax.legend();ax.grid(alpha=.15)
chart_save(fig,'training')
fig,axs=plt.subplots(1,3,figsize=(12.5,4),layout='constrained')
for c,ax in enumerate(axs):
 for method in ['ista','fista','lista']:
  for i in range(2):
   d=M[i]['config']['layers'];n=M[i]['config']['baseline_iterations']
   keys=['lista'] if method=='lista' else [f'{method}_{d}',f'{method}_{n}']
   for key in keys:
    entry=M[i]['comparisons'][key];xx=entry['inference']['median_seconds'];yy=100*entry['metrics'][c]['nrmse_range'];count=d if method=='lista' else int(key.split('_')[1]);ax.scatter(xx,yy,color=COL[method],s=40,marker='o' if i==0 else 's');ax.annotate(str(count),(xx,yy),xytext=(3,4),textcoords='offset points',fontsize=9,color=COL[method])
 ax.set_xscale('log');ax.set_xlabel('Batch latency (s), log scale');ax.set_title(['Bx','By','Bz'][c]);ax.grid(alpha=.15)
axs[0].set_ylabel('Range NRMSE (%)');chart_save(fig,'latency_tradeoff')
# Soft threshold visual.
fig,ax=plt.subplots(figsize=(5,2.5),layout='constrained');v=np.linspace(-2,2,400);ax.plot(v,v,'--',color='gray',alpha=.5,label='Identity');ax.plot(v,np.sign(v)*np.maximum(abs(v)-.5,0),color=COL['ista'],lw=3,label='Soft threshold, θ = 0.5');ax.axhline(0,color='gray',lw=.5);ax.axvline(0,color='gray',lw=.5);ax.set_xlabel('Input');ax.set_ylabel('Output');ax.legend(fontsize=9);chart_save(fig,'shrink')

R=Presentation();R.slide_width=Inches(13.333);R.slide_height=Inches(7.5)
slides=[]
def rgb(s):return RGBColor.from_string(s)
def box(s,x,y,w,h,fill=LIGHT,line=None):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
 if line:sh.line.color.rgb=rgb(line)
 else:sh.line.fill.background()
 return sh
def text(s,x,y,w,h,t,size=20,bold=False,color=DARK,font='Aptos',align=PP_ALIGN.LEFT):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for j,line in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.alignment=align;p.space_after=Pt(10)
 return sh
def slide(title,kicker='METHOD',notes=''):
 s=R.slides.add_slide(R.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(WHITE)
 box(s,0,0,13.333,.12,NAVY);text(s,.55,.3,12,.3,kicker,11,True,BLUE);text(s,.55,.82,12.2,.9,title,29,True,NAVY)
 text(s,.55,7.13,11.7,.22,'COMSOL z = 3 → 0.5 µm  •  Same-sample reconstruction; no independent test set',10,color=GRAY)
 text(s,12.15,7.1,.6,.25,str(len(R.slides)),11,color=GRAY,align=PP_ALIGN.RIGHT)
 s.notes_slide.notes_text_frame.text=notes;slides.append((title,notes));return s
def banner(s,t,color=BLUE):box(s,.55,6.38,12.2,.48,LIGHT);text(s,.72,6.46,11.85,.3,t,15,True,color)
def pic(s,p,x,y,w,h):
 iw,ih=Image.open(p).size;scale=min(w/iw,h/ih);ww=iw*scale;hh=ih*scale;s.shapes.add_picture(str(p),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))
def eq(s,formula,x=.7,y=2,w=11.9,h=1,size=28):
 fig=plt.figure(figsize=(12,1));fig.text(.5,.5,'$'+formula+'$',ha='center',va='center',fontsize=size,color='#'+NAVY);p=AS/f'eq_{len(list(AS.glob("eq_*.png"))):02d}.png';fig.savefig(p,dpi=220,transparent=True,bbox_inches='tight',pad_inches=.12);plt.close(fig);pic(s,p,x,y,w,h)
def bullets(s,items,x=.75,y=2,w=11.7,size=22,gap=.78):
 for i,item in enumerate(items):text(s,x,y+i*gap,w,.7,'•  '+item,size)
def table(s,headers,rows,x=.65,y=2,w=12,h=3.7,size=17,widths=None):
 sh=s.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(w),Inches(h));tb=sh.table
 if widths:
  for col,frac in zip(tb.columns,widths):col.width=Inches(w*frac)
 for i,row in enumerate([headers]+rows):
  for j,v in enumerate(row):
   c=tb.cell(i,j);c.fill.solid();c.fill.fore_color.rgb=rgb(NAVY if i==0 else (LIGHT if i%2 else WHITE));c.margin_left=Inches(.12);c.margin_right=Inches(.08);c.margin_top=Inches(.08);c.margin_bottom=Inches(.03)
   c.text=str(v)
   for p in c.text_frame.paragraphs:p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=(i==0);p.font.color.rgb=rgb(WHITE if i==0 else DARK)
 return tb

def image_row(s,filename,row,y=2.15):
 # PowerPoint cropping only: preserve the original PNG; select a component row.
 p=DATA/filename;iw,ih=Image.open(p).size
 labels=['Observed','Reference','ISTA 500' if '-2' in filename else 'ISTA 200','LISTA 16' if '-2' in filename else 'LISTA 8','FISTA 500' if '-2' in filename else 'FISTA 200']
 for j,label in enumerate(labels):text(s,.6+j*2.42,y-.27,2.42,.22,label,13,True,GRAY,align=PP_ALIGN.CENTER)
 top=[65,492,915][row];bottom=[475,905,1320][row]
 sh=s.shapes.add_picture(str(p),Inches(.6),Inches(y),width=Inches(12.1),height=Inches(12.1*(bottom-top)/iw))
 sh.crop_top=top/ih;sh.crop_bottom=max(0,1-bottom/ih)

s=slide('Recovering near-field magnetic structure','RESEARCH PROGRESS  /  ABHAY',"Present the goal as field-plane reconstruction, not current-density inversion. All reported observations and references are actual COMSOL exports. This deck reports two controlled, same-sample experiments. Main slides 1–20; technical backup follows.")
text(s,.7,2.0,8.2,1.2,'ISTA, FISTA and supervised\nconvolutional LISTA',34,True)
text(s,.7,3.65,7.6,1,'Actual COMSOL observations: z = 3 µm\nTarget field plane: z = 0.5 µm',23,color=GRAY)
box(s,9,2.15,3.5,2.8,NAVY);text(s,9.25,2.45,3,2.2,'2 controlled runs\n3 field components\n163 learned parameters',24,True,WHITE)
banner(s,'Question: what reconstruction quality can we obtain at a short inference budget?')

s=slide('Actual COMSOL fields define the input and target','PROBLEM',"Each component is a signed 1400×1400 field image with 0.1 µm pixel spacing. The observation is the z=3 cached COMSOL array, not an analytically propagated z=0.5 image. The target is the actual z=0.5 reference. Cropped display below shows the Bz row from the original first-run figure; the complete images appear in backup.")
image_row(s,'reconstructions.png',2,2.3)
text(s,.75,5.0,11.8,.85,'b: observed at 3 µm     s: reference at 0.5 µm     x: reconstructed estimate\n1400 × 1400 pixels per component  •  0.1 µm spacing  •  signed Bx, By, Bz',20)
banner(s,'The inverse problem estimates a field plane, not the coil current distribution.')

s=slide('Upward continuation suppresses fine spatial detail','FORWARD MODEL',"In a source-free homogeneous upper half-space with the decaying solution, each lateral Fourier mode decays exponentially with height. f is in cycles per micrometre. Reversing this attenuation amplifies discrepancies. This motivates regularization, but the experiment does not substitute the ideal exponential for its fitted transfer. See scripts/propagator.py for the analytical utility.")
eq(s,r'\widehat{B}(\mathbf{f},z+\Delta z)=e^{-2\pi|\mathbf{f}|\Delta z}\widehat{B}(\mathbf{f},z)',y=2.0,h=.95)
bullets(s,['Large spatial frequencies encode fine detail and decay more strongly.','Direct inversion can amplify numerical error and model discrepancy.','Our reconstruction uses a fitted discrete operator:  b ≈ Ax.'],y=3.35,size=22)
banner(s,'The ideal physical transfer motivates the problem; the experiment uses a COMSOL-fitted transfer.')

s=slide('One shared transfer is fitted from all three component pairs','FORWARD MODEL',"learn_kernel_joint in deblur.py minimizes the sum of squared Fourier residuals across Bx, By and Bz on the zero-padded grid. It is unregularized least squares with H set to zero where the absolute denominator is <=1e-12. H is fixed during LISTA training. The fit is not forced to be radial, positive or bounded by one. The reference is already used at this calibration stage, including for the baselines. Full padded fitting and cropped reconstruction losses have different domains.")
eq(s,r'H=\frac{\sum_c\overline{S_c}B_c}{\sum_c|S_c|^2},\qquad S_c=\mathcal{F}(Ps_c),\quad B_c=\mathcal{F}(Pb_c)',y=2,h=1.1)
bullets(s,['Joint fit across Bx, By and Bz; no cross-component mixing in inference.','H remains fixed while LISTA learns its correction filters and threshold.','No Wiener penalty or analytical-transfer constraint in this fit.'],y=3.45,size=21)
banner(s,'Reference-assisted calibration: the same image pairs are used for fitting and evaluation.',ORANGE)

s=slide('Padding and cropping are part of the operator','IMPLEMENTATION',"P zero-embeds the original image and C crops the same region, so C=P*. A=C T_H P, and the adjoint uses the conjugate transfer. Intermediate crop and re-padding in A* A are essential: replacing the composition by a single |H|² multiplier would generally change the operator. 32 pixels equals 3.2 µm per side. The padded FFT still implements circular convolution, with a zero-extension assumption; it does not establish an artifact-free physical boundary.")
for x,w,a,b in [(.7,3.6,'PAD','1400² → 1464²\n32 zeros on each side'),(4.8,3.6,'FFT OPERATOR','FFT → multiply by H\n→ inverse FFT → real'),(8.9,3.6,'CROP','1464² → 1400²\nRetain observed region')]:
 box(s,x,2.05,w,1.55,LIGHT);text(s,x+.18,2.22,w-.36,.4,a,17,True,BLUE);text(s,x+.18,2.75,w-.36,.75,b,19)
eq(s,r'A=CT_HP,\qquad A^\ast=CT_H^\ast P',y=4.0,h=.8)
text(s,.8,5.05,11.6,.8,'Adjoint: use conjugate H. Preserve the intermediate crop in A* A.\nZero padding reduces edge interaction; it does not guarantee its removal.',20)
banner(s,'“Retained padded operator” is precise; “exact physical model” would be too strong.')

s=slide('ISTA and FISTA solve the same pixel-domain L1 problem','OPTIMIZATION',"x is a signed field image. The first term is a sum of squared residuals, not a mean; the second is the sum of absolute field values. λ=10^-4 in microtesla units is inherited from the existing ISTA experiment, not selected by validation. L1 on pixels encourages small backgrounds but can suppress legitimate weak fields. There is no TV, wavelet, positivity, current prior or explicit Maxwell constraint. Treat this as a chosen convex objective, not a proof of physical truth.")
eq(s,r'\min_x J_\lambda(x)=\frac{1}{2}\|Ax-b\|_2^2+\lambda\|x\|_1',y=2.0,h=1.05)
box(s,.8,3.45,5.65,1.65,LIGHT);text(s,1,3.65,5.2,1.2,'Observation fit\nDoes the proposed sharp field reproduce b?',22)
box(s,6.8,3.45,5.65,1.65,LIGHT);text(s,7,3.65,5.2,1.2,'Pixel shrinkage\nPenalizes magnitude; encourages zeros.',22)
banner(s,'A smaller objective does not automatically mean a smaller error against the COMSOL reference.')

s=slide('ISTA corrects the residual, then soft-thresholds','ALGORITHMS',"For f(x)=0.5||Ax-b||², the gradient is A*(Ax-b). The proximal map of αλ||x||1 is soft thresholding. ISTA majorizes the smooth objective with a quadratic upper bound and minimizes that bound plus the L1 penalty. x0=b matches the repository initialization. The name is Iterative Shrinkage-Thresholding Algorithm. The norm bound and unit conventions are covered in backup.")
eq(s,r'x_{k+1}=\operatorname{soft}_{\alpha\lambda}\left(x_k-\alpha A^\ast(Ax_k-b)\right)',y=1.9,h=.9)
eq(s,r'\operatorname{soft}_{\theta}(v)=\operatorname{sign}(v)\max(|v|-\theta,0)',x=.65,y=3,w=7,h=.75,size=24)
text(s,.85,4.15,6.5,1.4,'x₀ = b\nα = 0.99 / max|H|² = 0.221052\nFixed threshold = αλ',23)
pic(s,AS/'shrink.png',7.7,3.1,4.8,2.7)
banner(s,'No training: every update follows the gradient and proximal map of the selected objective.')

s=slide('FISTA adds momentum to the same proximal update','ALGORITHMS',"Standard Beck–Teboulle FISTA uses y1=x0 and t1=1. The first step matches ISTA because momentum is initially zero. Return xk, not the extrapolated y. This implementation has no restart, backtracking or monotonicity modification. Under convex assumptions the objective-gap rate is O(1/k²), versus O(1/k) for ISTA. These are not guarantees on reference-image error or monotonic objective decrease at every step. Source: https://www.ceremade.dauphine.fr/~carlier/FISTA")
eq(s,r'x_k=\operatorname{soft}_{\alpha\lambda}\left(y_k-\alpha A^\ast(Ay_k-b)\right)',y=1.9,h=.8)
eq(s,r't_{k+1}=\frac{1+\sqrt{1+4t_k^2}}{2}',y=2.95,h=.85)
eq(s,r'y_{k+1}=x_k+\frac{t_k-1}{t_{k+1}}(x_k-x_{k-1})',y=4.0,h=.85)
text(s,.8,5.2,11.7,.65,'Same A, λ, α and initialization as ISTA. No learned parameters.',22)
banner(s,'Acceleration concerns the objective gap: O(1/k²) rather than O(1/k).')

s=slide('ConvLISTA learns corrections to a finite ISTA sequence','LEARNED RECONSTRUCTION',"This is a structured supervised LISTA variant, not a pure finite-filter replacement of the physics operator. Wb=αA*+Cb and Wx=I−αA*A+Cx. Both correction filters are bias-free, one input/output channel, stride one, 9×9, and tied across depth. θ is one scalar shared over space, components and layers. Input shape [3,1,1400,1400] treats components as batch samples, not feature channels. The whole layer retains global FFT interactions. No claim of convergence beyond trained depth. Source: scripts/lista.py.")
eq(s,r'x_{k+1}=\operatorname{soft}_{\theta}\left[x_k-\alpha A^\ast(Ax_k-b)+C_b(b)+C_x(x_k)\right]',y=1.85,h=.9,size=25)
table(s,['Fixed structure','Trainable correction'],[['Padded A and its adjoint','Cb: one 9 × 9 convolution'],['Base step α and starting image b','Cx: one 9 × 9 convolution'],['8 or 16 repeated layers','θ: one nonnegative scalar']],y=3,h=2.25,size=20)
text(s,.85,5.55,11.7,.5,'Tied parameters: 81 + 81 + 1 = 163 in both runs',25,True,ORANGE)
banner(s,'Initial Cb = Cx = 0 and θ = αλ reproduce ISTA exactly.')

s=slide('The supervised target changes what LISTA is trying to do','TRAINING OBJECTIVE',"Original LISTA learns to approximate sparse codes obtained by optimization. Here the user-selected target is the actual sharp COMSOL field. Only the final layer is supervised with full-image MSE. No additional data-consistency or L1 penalty is in the training loss. Therefore learned intermediate iterates need not reduce J, and learned maps need not be proximal gradient maps of any convex objective. Original paper: https://icml.cc/Conferences/2010/papers/449.pdf")
table(s,['Solver approximation','This experiment'],[['Target: converged L1 solution x*λ','Target: sharp COMSOL field s'],['Learn to imitate an optimizer','Learn final reference reconstruction'],['Solver output defines success','Reference MSE defines success']],y=1.95,h=2.4,size=20)
eq(s,r'\mathcal{L}(\Theta)=\frac{1}{3N}\sum_c\|F_\Theta(b_c)-s_c\|_2^2',y=4.65,h=.9)
banner(s,'The L1 objective is a diagnostic for LISTA, not its supervised training loss.',ORANGE)

s=slide('Only depth and baseline iteration budget changed','EXPERIMENT DESIGN',"The two JSON files have identical input and source hashes. Both CSV tables agree numerically with their JSON endpoint metrics. Run 1 uses 8 layers and 200 baseline iterations; run 2 uses 16 and 500. Both include matched-depth baselines. Training uses 200 full-batch Adam updates with lr=1e-4 and best training-loss selection. Same λ, pad, normalization, seed, hardware and code. Equal epochs do not imply equal training compute. Deterministic float32 GPU execution without AMP or TF32.")
table(s,['Setting','Run 1','Run 2'],[['LISTA depth / matched baseline steps','8 / 8','16 / 16'],['Long ISTA and FISTA budget','200','500'],['Training epochs / learning rate','200 / 10⁻⁴','200 / 10⁻⁴'],['Parameter count / filters','163 / tied 9 × 9','163 / tied 9 × 9'],['λ / FFT padding / initialization','10⁻⁴ / 32 / blurred','10⁻⁴ / 32 / blurred']],y=1.9,h=3.75,size=18,widths=[.48,.26,.26])
banner(s,'Same source + input hashes. Training: 2 × T4; inference comparisons: one T4.')

s=slide('At matched depth, LISTA has lower reference error','RESULTS  /  BOTH RUNS',"Every bar is taken from the saved JSON endpoint metrics, checked against CSVs. At both matched depths all three components have lower range NRMSE under trained LISTA. This comparison does not imply equal operation counts: LISTA adds convolutional corrections. Short-run batch inference is 0.216/0.211/0.240 s for ISTA8/FISTA8/LISTA8 and 0.418/0.431/0.460 s for ISTA16/FISTA16/LISTA16. These are same-sample supervised results, not held-out accuracy.")
pic(s,AS/'equal_depth.png',.65,1.8,12,4.3)
banner(s,'A LISTA layer costs more than a baseline step; its advantage here is error at a short budget.')

s=slide('More computation affects the three methods differently','RESULTS  /  DEPTH AND ITERATIONS',"Cross-run reference-error reductions: LISTA 3.46%, 2.66%, 3.44%; ISTA 6.38%, 5.43%, 9.28%. FISTA reference-error increases: 32.85%, 34.92%, 11.57%. The settings other than depth and iteration budget are controlled. This is one deterministic experiment per depth; no variability estimates or generalization claims. Baseline iteration count does not influence LISTA training.")
pic(s,AS/'between_runs.png',.65,1.8,12,4.25)
banner(s,'LISTA improves modestly; ISTA improves; FISTA endpoint reference error worsens.')

s=slide('FISTA reference error reaches a minimum before 500 steps','RESULTS  /  RUN 2 HISTORY',"The full run-2 history is available. FISTA's lowest recorded range NRMSE occurs at k=60 for Bx (1.616636%) and By (1.460708%), and k=250 for Bz (1.795615%). Metrics are sampled every ten iterations, so these are best recorded points, not exact minima. Selecting an iteration using the reference is an oracle selection; deployment requires separate validation or a justified observation-only rule. ISTA's best recorded points are at the 500-step endpoint in all components. The LISTA path ends at trained depth 16.")
pic(s,AS/'reference_history.png',.65,1.8,12,4.25)
banner(s,'Best recorded FISTA: Bx 60, By 60, Bz 250. Reference-based selection is not a deployable rule.')

s=slide('Optimization improves even when reference agreement worsens','RESULTS  /  TWO DIFFERENT CRITERIA',"For Bx, FISTA200→500 reduces J from 16576.105 to 5691.415 and PG RMS from 0.00327225 to 0.000637219, while range NRMSE increases from 1.86991% to 2.48416%. Analogous endpoint worsening occurs in By and Bz. This is evidence of objective/reference mismatch, not algorithm divergence. Semi-convergence is a useful description of the observed reference-error path; the physical cause of discrepancies is not isolated. LISTA is trained against a different loss.")
pic(s,AS/'optimization_history.png',.65,1.8,12,3.65)
text(s,.8,5.55,11.8,.6,'Bx FISTA 200 → 500: objective 16,576 → 5,691; NRMSE 1.870% → 2.484%',20,True)
banner(s,'A better optimum of the chosen model need not be a better estimate of the reference field.',ORANGE)

s=slide('Sixteen layers improve training MSE at twice the training cost','RESULTS  /  TRAINING',"Run 1 best training epoch=200, loss=0.001674862229; run 2 best epoch=200, loss=0.001565922867. Training durations include post-update evaluation: 150.1225 and 309.3839 seconds. Final physical thresholds θq are 0.746789 and approximately 0.31694 µT. Their difference is not an interpretable lambda change because the learned linear maps also change. Both use one shared nonnegative threshold. Each curve is final-output loss across epochs, not convergence across layers.")
pic(s,AS/'training.png',.65,1.8,12,3.75)
text(s,.8,5.65,11.8,.45,'Best epoch: 200 in both runs  •  Training: 150.1 s → 309.4 s  •  Parameters: 163 → 163',20,True)
banner(s,'Equal epoch counts do not equalize training compute; no validation set was used.')

s=slide('The useful comparison is reconstruction error versus cost','RESULTS  /  INFERENCE',"All endpoint latencies are warmed synchronized medians of three repeats, on the same single GPU and batch of three, excluding metrics, file I/O, kernel fitting and training. Color indicates algorithm, marker indicates run, labels give iterations/layers. LISTA8 0.239503 s; LISTA16 0.459688 s. ISTA200/500 5.180782/13.073968 s; FISTA200/500 5.303321/13.485866 s. Long baselines have different accuracy, so no equal-accuracy 20x/30x speedup is claimed. Three timing repeats do not provide a broad benchmark distribution.")
pic(s,AS/'latency_tradeoff.png',.65,1.8,12,4.1)
text(s,.85,6.0,11.7,.28,'Blue: ISTA   •   Green: FISTA   •   Orange: LISTA   |   Labels: steps/layers; circles: run 1, squares: run 2',13)
banner(s,'LISTA is fast at the tested depth; longer baselines can achieve lower reference error.')

s=slide('Run 2: sharper structure and background artifacts must both be assessed','RESULTS  /  Bz VISUAL COMPARISON',"This is the Bz row of reconstructions-2.png, cropped by PowerPoint without modifying the source image. Columns are observed, reference, ISTA500, LISTA16, FISTA500. Shared row display limits are ±99.5th percentile of |reference|, so extreme values saturate; there is no independent autoscaling per reconstruction. Smoothness is not sufficient evidence of fidelity. FISTA500 has lower Bz endpoint error than ISTA500, and both improve on LISTA16. FISTA200 has lower Bz endpoint error than FISTA500. The speckle's cause has not been isolated.")
image_row(s,'reconstructions-2.png',2,2.0)
text(s,.8,4.7,11.8,.9,'Bz range NRMSE: ISTA 500 = 2.145%   |   LISTA 16 = 2.861%   |   FISTA 500 = 2.026%\nShared color range within the row; saturated extremes; full field of view.',20)
banner(s,'A smoother-looking output is not automatically closer to the reference.')

s=slide('The evidence supports a tradeoff, not a universal winner','INTERPRETATION',"Among the endpoint rows in the two CSVs, ISTA500 is best for Bx and By; FISTA200 is best for Bz. Across saved histories, FISTA250 slightly improves Bz further, but that is reference-based oracle selection, not a validated configuration. LISTA wins reference error at matched depth in both runs. FISTA500 gives the smallest objective and stationarity residual among endpoints. Range NRMSE divides by full reference range and is not percent accuracy. LISTA16 relative L2 errors are 42.15/43.95/48.44%.")
table(s,['Criterion','Supported result'],[['Reference error at matched 8/16 updates','LISTA is lower in all three components'],['Best Bx / By endpoint in the two CSVs','ISTA 500: 1.595% / 1.443%'],['Best Bz endpoint in the two CSVs','FISTA 200: 1.816%'],['Smallest original objective among endpoints','FISTA 500'],['Generalization to unseen observations','Not evaluated']],y=1.95,h=3.9,size=19,widths=[.49,.51])
banner(s,'Range NRMSE ≈ 2% does not mean “98% accurate”; also report relative L2 and physical RMSE.')

s=slide('Next: separate the objective, calibration and evaluation questions','DISCUSSION',"Proposed follow-up work, not completed experiments. First clarify whether the research goal is solver acceleration or physical reference reconstruction. Use separate validation data for λ and stopping/depth selection, and independent test observations for final comparisons. Assess forward residuals and boundaries separately; consider analytical or independently calibrated transfers if justified. Compare at matched accuracy or latency and count training costs. Repeat over seeds/data to assess robustness. Main presentation ends here; backup slides follow.")
bullets(s,['Choose the target: approximate the L1 optimizer or recover a sharp reference.','Validate λ and stopping rules separately from final evaluation.','Test independently calibrated models on unseen observations.','Compare matched accuracy / latency, with training cost reported.'],y=2,size=23,gap=.93)
banner(s,'Current result: promising short-budget reconstruction, demonstrated only on the training sample.')

# Technical backup.
s=slide('Backup: derive shrinkage from the proximal subproblem','TECHNICAL BACKUP',"Start with f(x)=0.5||Ax-b||². Gradient is A*(Ax-b). The quadratic upper bound with step α leads to minimizing 0.5||u-v||²+αλ||u||1, where v=x−α∇f(x). The problem separates over pixels. For positive u, derivative gives u=v−θ; for negative u, u=v+θ; if |v|≤θ the solution is zero. This explains soft thresholding and why signed field values are preserved. L1 penalizes amplitudes; it is not an edge-preserving TV penalty.")
eq(s,r'v=x-\alpha\nabla f(x),\qquad \nabla f(x)=A^\ast(Ax-b)',y=1.9,h=.8)
eq(s,r'\operatorname{prox}_{\alpha\lambda\|\cdot\|_1}(v)=\arg\min_u\left\{\frac{1}{2}\|u-v\|^2+\alpha\lambda\|u\|_1\right\}',y=3,h=1,size=25)
eq(s,r'u_i=\operatorname{sign}(v_i)\max(|v_i|-\alpha\lambda,0)',y=4.4,h=.85)
banner(s,'For fixed A and λ ≥ 0, the baseline objective is convex.')

s=slide('Backup: norm bound, units and initial threshold','TECHNICAL BACKUP',"C and P have operator norm one; padded circular convolution has norm max|H|. Thus Ltrue=||A||²≤Lbound=max|H|²=4.478589576. α=0.99/Lbound=0.221051736. With x=q x_tilde and q=215.1441345 µT, J(q x_tilde)=q²[0.5||A x_tilde−b_tilde||²+(λ/q)||x_tilde||1]. Normalized λ=4.64804677e-7 and initial θ=1.02745879e-7. Physical fixed threshold=2.21051736e-5 µT. The shared scale uses both reference and observed images; saved inference must reuse the checkpoint scale.")
eq(s,r'\|A\|_2^2\leq\max|H|^2=L_{\rm bound},\qquad\alpha=0.99/L_{\rm bound}',y=1.9,h=.8,size=26)
eq(s,r'x=q\widetilde{x}\ \Longrightarrow\ J(x)=q^2\left[\frac{1}{2}\|A\widetilde{x}-\widetilde{b}\|^2+\frac{\lambda}{q}\|\widetilde{x}\|_1\right]',y=3,h=1,size=25)
table(s,['Quantity','Both runs'],[['q / normalized λ','215.1441 µT / 4.64805 × 10⁻⁷'],['Lbound / α','4.47859 / 0.221052'],['Initial threshold in normalized units','1.02746 × 10⁻⁷']],y=4.4,h=1.65,size=17)

s=slide('Backup: training and tensor implementation','TECHNICAL BACKUP',"DataParallel splits batch size three into 2+1. MSE is computed after gathering predictions, preventing unequal chunk weighting. Model parameters are only the two filters and threshold; H real and imaginary arrays are nontrainable buffers. Direct θ parameter is projected nonnegative after Adam; clamp_min also protects forward evaluation. Activation checkpointing recomputes layers during backpropagation to reduce memory. Best weights are selected after post-update reevaluation. Kernel fitted in float64, GPU inference float32; no AMP or TF32. These details do not establish generalization.")
table(s,['Detail','Implementation'],[['Input / component handling','[3, 1, 1400, 1400]; shared weights, no channel mixing'],['Correction convolutions','9 × 9, stride 1, no bias; 4-pixel zero padding'],['Parameter sharing','Two filters + one θ reused at every layer'],['Optimizer','Adam lr 10⁻⁴; gradient-norm clipping at 1'],['Selection / memory','Best post-update training MSE; activation checkpointing'],['Numeric execution','Float32 GPU; fixed H buffers; no AMP / TF32']],y=1.85,h=4.2,size=17,widths=[.32,.68])

s=slide('Backup: the metrics answer different questions','TECHNICAL BACKUP',"Reference RMSE is in microtesla. Range NRMSE normalizes by max(s)−min(s); relative L2 normalizes by ||s||. The denominator difference explains seemingly small range NRMSE despite sizeable relative L2. PG is the proximal-gradient mapping of the original objective, not the training gradient of the learned network. PG=0 is optimality for the convex baseline objective, but small PG alone does not bound reference error or distance to a unique optimizer in an ill-conditioned system. All field metrics use the full field of view.")
eq(s,r'\mathrm{RMSE}=\frac{\|x-s\|_2}{\sqrt{N}},\quad\mathrm{NRMSE}_{range}=\frac{\mathrm{RMSE}}{\max(s)-\min(s)},\quad\mathrm{RelL2}=\frac{\|x-s\|_2}{\|s\|_2}',y=2,h=1,size=24)
eq(s,r'G_\alpha(x)=\frac{x-\operatorname{soft}_{\alpha\lambda}(x-\alpha\nabla f(x))}{\alpha}',y=3.45,h=1)
text(s,.8,5,11.7,1,'Stationarity RMS = ||Gα(x)||₂ / √N\nLISTA 16 relative L2: 42.15% (Bx), 43.95% (By), 48.44% (Bz)',23)

s=slide('Backup: map the mathematics to source files','CODE MAP',"Notebook extracts and verifies the bundle, runs tests, launches scripts/lista.py and displays/export results. lista.py CLI delegates to compare_lista.py, while its ConvLISTA class is independently importable. compare_lista orchestrates fitting, scaling, baseline solving, training, metrics and benchmarks. GPU baselines use torch_inverse.py; standalone CPU FISTA uses iterative.py. Original deblur.py remains the fit/operator reference. Loading the checkpoint includes saved H, scale and original norm bound to preserve inference behavior.")
table(s,['File','Responsibility'],[['compare_lista.py','Data, fit, normalization, training, metrics, timing, export'],['lista.py','ConvLISTA recurrence, tied parameters, checkpoint loading'],['torch_inverse.py','Batched padded A / A*, GPU ISTA/FISTA and metrics'],['deblur.py','Original fit, NumPy operators and ISTA'],['iterative.py / fista.py','Shared NumPy solvers / standalone FISTA command'],['plot_comparison.py / prepare_kaggle.py','Saved-result figures / notebook and portable bundle']],y=1.85,h=4.3,size=15,widths=[.35,.65])

s=slide('Backup: checks support implementation, not physical optimality','VALIDATION AND ASSUMPTIONS',"Tests in test_iterative.py and test_lista.py cover adjoint identities, dense references, known L1 solution, ISTA initialization parity, finite-difference gradients, normalization, checkpointing and uneven two-GPU gradients. The first local run previously passed 12 with the two-GPU test skipped; notebook includes that check, but these result files do not store its test log. Therefore this slide describes coverage and run-time initialization result, without inventing a GPU test count. Both run JSONs report initialization max error 0. No independent samples or uncertainty estimates.")
table(s,['Implementation checks','Modeling / evaluation limits'],[['Adjoint and dense reference updates','Finite FOV and zero-extension boundaries'],['Known L1 solution and ISTA parity','Shared fitted transfer; pixel-domain L1 prior'],['Gradients, normalization, checkpoint reload','Only one coil and one height pair'],['Uneven two-GPU gradient test included','Reference used in fitting, scale and training'],['Initialization max difference = 0 in both runs','No validation set, test set or seed variability']],y=1.9,h=3.8,size=18)
banner(s,'No conclusion of generalization, exact physics, or optimal hyperparameters follows from these checks.')

for i in range(2):
 s=slide(f'Backup: run {i+1} endpoint metrics','NUMERICAL RESULTS',"Numbers are read directly from the corresponding metrics JSON; CSVs were checked for agreement. Latency is a median for the batch of all three components, not per component. Range NRMSE values are percentages; JSON stores fractions. All comparisons are same-sample. The listed best endpoints must not be confused with minima selected over histories.")
 rows=[]
 for key,entry in M[i]['comparisons'].items():
  label=f'LISTA {M[i]["config"]["layers"]}' if key=='lista' else key.replace('_',' ').upper()
  rows.append([label]+[f'{v:.3f}' for v in 100*scores(i,key)]+[f'{entry["inference"]["median_seconds"]:.3f}'])
 table(s,['Method','Bx NRMSE %','By NRMSE %','Bz NRMSE %','Time (s)'],rows,y=2,h=3.6,size=19,widths=[.26,.19,.19,.19,.17])
 banner(s,'Short-depth comparisons and long-run endpoints serve different comparison questions.')
for i in range(2):
 s=slide(f'Backup: complete reconstruction grid — run {i+1}','VISUAL RESULTS',"Original supplied figure, preserved without pixel editing. Columns compare observed, reference, long ISTA, trained LISTA, long FISTA. Shared per-row display range is ±99.5th percentile of absolute reference; extremes can saturate. The same-sample label and unequal iteration budgets are explicit. Numerical arrays were not provided in this results folder, so no new physical ROI metrics or residual maps were synthesized.")
 pic(s,DATA/('reconstructions.png' if i==0 else 'reconstructions-2.png'),.65,1.7,12,5.2)

s=slide('Backup: provenance and primary references','REFERENCES',"Results are supplied in notebooks/runs_results/metrics.json and metrics-2.json, with matching source/input SHA256 dictionaries. CSVs numerically agree with the JSONs. Source references are Gregor & LeCun (2010), Learning Fast Approximations of Sparse Coding, ICML; Beck & Teboulle (2009), A Fast Iterative Shrinkage-Thresholding Algorithm for Linear Inverse Problems, SIAM J Imaging Sciences; PyTorch 2.10 Conv2d and Adam documentation. No new reconstruction or training was run to create the slides; plots are regenerated from saved metrics. Build script: presentations/build_professor_deck.py.")
text(s,.8,1.95,11.7,1.0,'Gregor & LeCun (2010) — Learning Fast Approximations of Sparse Coding\nhttps://icml.cc/Conferences/2010/papers/449.pdf',19)
text(s,.8,3.15,11.7,1.0,'Beck & Teboulle (2009) — A Fast Iterative Shrinkage-Thresholding Algorithm\nhttps://www.ceremade.dauphine.fr/~carlier/FISTA',19)
text(s,.8,4.35,11.7,1.0,'PyTorch 2.10 — Conv2d and Adam implementation documentation\nhttps://docs.pytorch.org/docs/2.10/',19)
text(s,.8,5.6,11.7,.5,'Data: notebooks/runs_results/  •  Identical input and source hashes across runs',18,True)

R.core_properties.title='ISTA, FISTA and supervised ConvLISTA: COMSOL reconstruction'
R.core_properties.subject='Two controlled same-sample experiments, z=3 to 0.5 micrometres'
R.core_properties.author='Abhay'
R.save(OUT/'ISTA_FISTA_LISTA_Professor.pptx')
(OUT/'speaker_notes.md').write_text('\n\n'.join(f'{i}. {t}\n\n{n}' for i,(t,n) in enumerate(slides,1)))
manifest={'slides':len(slides),'main_slides':20,'results':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in DATA.iterdir() if p.is_file()},'source_hashes_match':True,'input_hashes_match':True}
(OUT/'presentation_manifest.json').write_text(json.dumps(manifest,indent=2))
print(f'Created {len(slides)} slides with speaker notes: {OUT}/ISTA_FISTA_LISTA_Professor.pptx')
